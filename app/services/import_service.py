"""Import wizard backend (Konzept §12): upload → preview → validate → commit → undo.

Flow
----
1. :func:`upload` stores the original file under ``DATA_DIR/imports/<batch-id>/original.<csv|xlsx>``,
   creates an ``import_batch`` row (stage ``hochgeladen``) and returns a preview: detected type, encoding,
   delimiter, sheets, header row, the first 50 data rows and a guessed column mapping.
   CSV: encoding utf-8 / utf-8-sig / utf-16 (BOM) / cp1252, delimiter sniffed among ``;``, tab and ``,``.
   XLSX: openpyxl (read-only, cached values), all worksheets selectable.
2. :func:`validate` parses every row with ``app.core.parsing`` (date, time, duration incl. ``2,5`` and
   ``h:mm``), resolves categories (free text → catalogue code: user map for this import, then the
   remembered ``category_mapping`` table, then catalogue code/name, then the default category), and
   reports per-row errors/warnings (a bounded row whose Bis lies before Von is an error unless Bis is
   written as next day – ``01:30+1`` – or a Dauer column confirms the overnight span; ``24:00`` is the end of
   the day; rows dated outside every period get a warning and are not imported unless
   ``include_outside_periods`` or ``include_rows`` says so), duplicates (same day + same duration + same main category; bounded rows
   additionally need the same start when compared with another bounded row) within the file and against
   existing entries, overlaps of bounded rows, and monthly sums file vs. import.
3. :func:`commit` re-runs the analysis and creates the entries through ``app.services.entries`` with
   ``source='import'`` and ``import_batch_id`` (rows with only a duration → ``has_bounds=False`` →
   flag ``ohne_zeitgrenzen``). Atomic: one failing row rolls back everything. User category mappings are
   remembered (Merkfunktion). Bounded imported rows are checked for overlaps here because the entries
   service exempts ``source='import'`` from its own overlap check.
4. :func:`undo` soft-deletes all entries of a batch (reason ``Import rückgängig``) through the entries
   service. Entries in closed months need a reason, otherwise the whole undo is refused
   (:class:`ImportLocked` with the count). A batch that was never committed is simply discarded.

Batch state lives in ``import_batch.mapping_json`` (no schema change)::

    {"stage": "hochgeladen"|"geprueft"|"importiert"|"rueckgaengig", "file_type": "csv"|"xlsx",
     "encoding", "delimiter", "sheets": [...], "size_bytes", "mapping": {field: column}, "selected_sheets",
     "header_row", "options": {...}, "committed_at", "entries_created", "counts": {...}}

Mapping: ``{field: column}`` with fields :data:`FIELDS`; a column is a 0-based index, a header text or a
column letter (``"B"``). Required: ``datum`` plus ``dauer`` or ``von``+``bis`` (``von``+``dauer`` and
``bis``+``dauer`` work too); ``kategorie`` or ``options.default_category``.

Row references (``ref``): the file row number (``"12"``) for CSV, ``"<sheet>!<row>"`` for XLSX.
``include_rows``/``exclude_rows`` accept refs or plain row numbers.
"""

from __future__ import annotations

import codecs
import csv
import io
import re
import shutil
from collections import Counter, defaultdict
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from itertools import islice
from pathlib import Path
from typing import Any

from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.config import get_config
from app.core.parsing import parse_date, parse_duration, parse_time_seconds
from app.db.models import (
    Category,
    CategoryMapping,
    EntrySource,
    EntryStatus,
    ImportBatch,
    Location,
    MonthClose,
    Period,
    Settings,
    TimeEntry,
)
from app.services import entries as entries_svc
from app.util import (
    fmt_hm,
    local_date,
    local_midnight_utc,
    local_to_utc,
    now_utc,
    to_local,
    to_utc_naive,
    today_local,
    uuid7,
)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLS = 100
PREVIEW_ROWS = 50
HEADER_SCAN_ROWS = 20
SNIFF_RECORDS = 30
DURATION_TOLERANCE_SECONDS = 59
MAX_DAY_SECONDS = 24 * 3600
EXCEL_EPOCH_DT = datetime(1899, 12, 30)

UNDO_REASON = "Import rückgängig"
COMMIT_REASON_PREFIX = "Import"

STAGE_UPLOADED = "hochgeladen"
STAGE_VALIDATED = "geprueft"
STAGE_COMMITTED = "importiert"
STAGE_UNDONE = "rueckgaengig"

STATUS_OK = "ok"
STATUS_ERROR = "fehler"
STATUS_DUPLICATE = "duplikat"
STATUS_EMPTY = "leer"
STATUS_IGNORED = "ignoriert"

FIELDS = ("datum", "von", "bis", "dauer", "kategorie", "klasse", "fach", "notiz", "ort")
TIME_FIELDS = ("datum", "von", "bis", "dauer")
FIELD_LABELS = {
    "datum": "Datum",
    "von": "Von",
    "bis": "Bis",
    "dauer": "Dauer",
    "kategorie": "Kategorie",
    "klasse": "Klasse",
    "fach": "Fach",
    "notiz": "Notiz",
    "ort": "Ort",
}

# Normalised header synonyms (see _norm_header). Exact match beats "starts with" beats "contains word".
HEADER_SYNONYMS: dict[str, tuple[str, ...]] = {
    "datum": ("datum", "date", "tag", "arbeitstag", "wochentag datum"),
    "von": ("von", "beginn", "start", "anfang", "startzeit", "uhrzeit von", "zeit von", "begin", "kommen"),
    "bis": ("bis", "ende", "end", "endzeit", "schluss", "uhrzeit bis", "zeit bis", "gehen"),
    "dauer": ("dauer", "stunden", "std", "zeit", "arbeitszeit", "duration", "stundenzahl", "zeitaufwand",
              "aufwand", "h"),
    "kategorie": ("kategorie", "taetigkeit", "taetigkeiten", "art", "bereich", "category", "aufgabe",
                  "taetigkeitsbereich", "arbeitsbereich"),
    "klasse": ("klasse", "klassen", "class"),
    "fach": ("fach", "faecher", "subject"),
    "notiz": ("notiz", "notizen", "bemerkung", "bemerkungen", "kommentar", "beschreibung", "anmerkung",
              "note", "inhalt", "details"),
    "ort": ("ort", "arbeitsort", "location"),
}

LOCATION_SYNONYMS = {
    "schule": Location.SCHULE.value,
    "in der schule": Location.SCHULE.value,
    "schulhaus": Location.SCHULE.value,
    "zuhause": Location.ZUHAUSE.value,
    "zu hause": Location.ZUHAUSE.value,
    "daheim": Location.ZUHAUSE.value,
    "home": Location.ZUHAUSE.value,
    "homeoffice": Location.ZUHAUSE.value,
    "home office": Location.ZUHAUSE.value,
    "unterwegs": Location.UNTERWEGS.value,
    "extern": Location.EXTERN.value,
    "auswaerts": Location.EXTERN.value,
}

_SUM_WORDS_RE = re.compile(
    r"\b(summe|summen|gesamt|gesamtsumme|total|uebertrag|zwischensumme|wochensumme|monatssumme|"
    r"jahressumme|insgesamt)\b"
)
_WEEKDAY_PREFIX_RE = re.compile(r"^[A-Za-zÄÖÜäöü]{2,10}\.?,?\s+(?=\d)")
_COLUMN_LETTER_RE = re.compile(r"^[A-Za-z]{1,2}$")
# Explicit next-day notation for "Bis": "01:30+1", "01:30 (+1)", "01:30 +1d", "01:30 +1 Tag".
_NEXT_DAY_RE = re.compile(r"\s*\(?\s*\+\s*1\s*(?:d|t|tag)?\s*\)?\s*$", re.IGNORECASE)

OPTION_KEYS = frozenset({
    "default_category", "category_map", "remember_mappings", "skip_errors", "skip_duplicates",
    "include_rows", "exclude_rows", "reason", "include_outside_periods",
})

OUTSIDE_PERIOD_WARNING = "Datum liegt außerhalb jedes Schuljahres"


# --------------------------------------------------------------------------- errors


class ImportServiceError(Exception):
    """Base class; ``status_code`` is the HTTP status the API uses, ``message`` is German."""

    status_code = 422

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


class ImportValidationError(ImportServiceError):
    status_code = 422


class ImportFileError(ImportServiceError):
    """Unreadable / unsupported file."""

    status_code = 422


class ImportFileTooLarge(ImportFileError):
    status_code = 413


class ImportNotFound(ImportServiceError):
    status_code = 404


class ImportConflict(ImportServiceError):
    """Batch already committed / undone."""

    status_code = 409


class ImportLocked(ImportServiceError):
    """Undo touches closed months and no reason was given."""

    status_code = 423

    def __init__(self, message: str, *, locked_entries: int = 0):
        super().__init__(message)
        self.locked_entries = locked_entries


# --------------------------------------------------------------------------- small helpers


def _umlauts(text: str) -> str:
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        text = text.replace(a, b)
    return text


def _norm_key(value: Any) -> str:
    """Normalisation for category texts: case-insensitive, umlaut-insensitive, collapsed whitespace."""
    return " ".join(_umlauts(str(value).casefold()).split())


def _norm_header(value: Any) -> str:
    """Normalisation for header names: like :func:`_norm_key`, punctuation becomes a space."""
    return " ".join(re.sub(r"[^0-9a-z]+", " ", _umlauts(str(value).casefold())).split())


def column_letter(index: int) -> str:
    """0 → 'A', 25 → 'Z', 26 → 'AA'."""
    result = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def _letter_index(text: str) -> int:
    value = 0
    for ch in text.upper():
        value = value * 26 + (ord(ch) - 64)
    return value - 1


def _clean_cell(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _cell_text(value: Any) -> str | None:
    """Plain text of a cell for text fields (Klasse, Fach, Notiz ...)."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime, date, time, timedelta)):
        value = display_cell(value)
    text = str(value).strip()
    return text or None


def display_cell(value: Any) -> str:
    """Cell value as shown in the preview (German formats)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "WAHR" if value else "FALSCH"
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.strftime("%d.%m.%Y")
        return value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S" if value.second else "%H:%M")
    if isinstance(value, timedelta):
        return fmt_hm(value.total_seconds())
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.10g}".replace(".", ",")
    return str(value)


def _iso_utc(dt: datetime | None) -> str | None:
    return None if dt is None else dt.isoformat() + "Z"


def _ym(d: date) -> str:
    return f"{d.year:04d}-{d.month:02d}"


MONTH_NAMES = ("Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September",
               "Oktober", "November", "Dezember")


def _month_label(ym: str) -> str:
    y, m = ym.split("-")
    return f"{MONTH_NAMES[int(m) - 1]} {y}"


def _local_dt(day: date, seconds: int) -> datetime:
    """Local wall-clock ``day`` + ``seconds`` since midnight (86400 = next midnight) → naive UTC."""
    if seconds >= MAX_DAY_SECONDS:
        return local_midnight_utc(day + timedelta(days=1))
    return local_to_utc(day, time(seconds // 3600, seconds % 3600 // 60, seconds % 60))


def to_naive(dt: datetime) -> datetime:
    """Naive UTC without microseconds (DB format); naive input is UTC."""
    return to_utc_naive(dt).replace(microsecond=0)


def _hm_local(dt: datetime | None) -> str | None:
    return None if dt is None else to_local(dt).strftime("%H:%M")


# --------------------------------------------------------------------------- file reading


def detect_file_type(content: bytes, filename: str | None = None) -> str:
    """'xlsx' or 'csv' (by content, the extension is only used for error messages)."""
    if content.startswith(b"PK\x03\x04"):
        return "xlsx"
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        raise ImportFileError(
            "Das alte Excel-Format (.xls) wird nicht unterstützt. Bitte als .xlsx oder .csv speichern."
        )
    if (filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise ImportFileError("Die Datei ist keine gültige XLSX-Datei.")
    return "csv"


def decode_text(content: bytes) -> tuple[str, str]:
    """Decode CSV bytes; returns (text, encoding). BOM → utf-8-sig / utf-16, else utf-8, else cp1252."""
    if content.startswith(codecs.BOM_UTF8):
        try:
            return content[len(codecs.BOM_UTF8):].decode("utf-8"), "utf-8-sig"
        except UnicodeDecodeError as exc:
            raise ImportFileError("Die Datei ist nicht korrekt UTF-8-kodiert.") from exc
    if content.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        try:
            return content.decode("utf-16"), "utf-16"
        except UnicodeDecodeError as exc:
            raise ImportFileError("Die Datei ist nicht korrekt UTF-16-kodiert.") from exc
    try:
        text, encoding = content.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        text, encoding = content.decode("cp1252", errors="replace"), "cp1252"
    if "\x00" in text:
        raise ImportFileError("Die Datei ist weder CSV noch XLSX.")
    return text, encoding


def sniff_delimiter(text: str) -> str:
    """Pick the delimiter among ``;``, tab and ``,`` giving the most consistent multi-column rows.

    Ties prefer ``;`` (German Excel), then tab, then ``,`` – so decimal commas never win against ``;``.
    """
    best: tuple[tuple[int, float, int], str] | None = None
    for delim in (";", "\t", ","):
        try:
            counts = [len(r) for r in islice(csv.reader(io.StringIO(text), delimiter=delim), SNIFF_RECORDS)
                      if any(c.strip() for c in r)]
        except csv.Error:
            continue
        if not counts:
            continue
        mode, freq = Counter(counts).most_common(1)[0]
        score = (1, freq / len(counts), mode) if mode >= 2 else (0, 0.0, 0)
        if best is None or score > best[0]:
            best = (score, delim)
    return best[1] if best is not None and best[0][0] else ";"


def _read_csv(content: bytes, encoding: str | None = None, delimiter: str | None = None
              ) -> tuple[list[tuple[int, list[Any]]], str, str]:
    if encoding:
        try:
            text = content.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            text, encoding = decode_text(content)
        else:
            text = text.lstrip("﻿")
    else:
        text, encoding = decode_text(content)
    delimiter = delimiter or sniff_delimiter(text)
    rows: list[tuple[int, list[Any]]] = []
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        for record in reader:
            cells = [_clean_cell(c) for c in record[:MAX_COLS]]
            if any(c is not None for c in cells):
                rows.append((reader.line_num, cells))
                if len(rows) > MAX_ROWS:
                    raise ImportFileError(f"Die Datei hat zu viele Zeilen (höchstens {MAX_ROWS:,}).".replace(",", "."))
    except csv.Error as exc:
        raise ImportFileError(f"CSV-Datei konnte nicht gelesen werden: {exc}") from exc
    return rows, encoding, delimiter


def _open_workbook(content: bytes):
    from openpyxl import load_workbook

    try:
        return load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds (BadZipFile, KeyError, InvalidFileException ...)
        raise ImportFileError("Die Datei ist keine gültige XLSX-Datei.") from exc


def _xlsx_sheet_names(content: bytes) -> list[str]:
    wb = _open_workbook(content)
    try:
        names = [ws.title for ws in wb.worksheets]
    finally:
        wb.close()
    if not names:
        raise ImportFileError("Die Arbeitsmappe enthält keine Tabellenblätter.")
    return names


def _read_xlsx(content: bytes, sheet: str) -> list[tuple[int, list[Any]]]:
    wb = _open_workbook(content)
    try:
        if sheet not in [ws.title for ws in wb.worksheets]:
            raise ImportValidationError(f"Tabellenblatt „{sheet}“ nicht gefunden.")
        ws = wb[sheet]
        rows: list[tuple[int, list[Any]]] = []
        for row_no, values in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [_clean_cell(v) for v in values[:MAX_COLS]]
            while cells and cells[-1] is None:
                cells.pop()
            if cells:
                rows.append((row_no, cells))
                if len(rows) > MAX_ROWS:
                    raise ImportFileError(f"Das Blatt hat zu viele Zeilen (höchstens {MAX_ROWS:,}).".replace(",", "."))
        return rows
    except ImportServiceError:
        raise
    except Exception as exc:
        raise ImportFileError(f"Tabellenblatt „{sheet}“ konnte nicht gelesen werden.") from exc
    finally:
        wb.close()


def _header_score(header: str, field_name: str) -> int:
    norm = _norm_header(header)
    if not norm:
        return 0
    best = 0
    for syn in HEADER_SYNONYMS[field_name]:
        if norm == syn:
            return 3
        if norm.startswith(syn + " "):
            best = max(best, 2)
        elif len(syn) > 2 and f" {syn} " in f" {norm} ":
            best = max(best, 1)
    return best


def _looks_like_date(value: Any) -> bool:
    if value is None or isinstance(value, (bool, time, timedelta)):
        return False
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return False  # bare numbers are ambiguous (Excel serials, durations)
    try:
        _parse_day(value)
        return True
    except ValueError:
        return False


def _is_header_row(cells: list[Any]) -> bool:
    texts = [c for c in cells if isinstance(c, str)]
    if any(_looks_like_date(c) for c in cells):
        return False
    hits = sum(1 for t in texts if any(_header_score(t, f) for f in FIELDS))
    return hits >= 2


@dataclass
class _Table:
    sheet: str | None
    rows: list[tuple[int, list[Any]]]  # non-empty rows: (1-based row number in the file, cells)
    header_index: int | None  # index into ``rows``; None = no header row

    @property
    def ncols(self) -> int:
        return max((len(c) for _, c in self.rows), default=0)

    @property
    def header_row(self) -> int:
        return self.rows[self.header_index][0] if self.header_index is not None else 0

    @property
    def headers(self) -> list[str]:
        cells = self.rows[self.header_index][1] if self.header_index is not None else []
        result = []
        for i in range(self.ncols):
            value = cells[i] if i < len(cells) else None
            text = display_cell(value).strip() if value is not None else ""
            result.append(text or f"Spalte {column_letter(i)}")
        return result

    @property
    def data(self) -> list[tuple[int, list[Any]]]:
        return self.rows if self.header_index is None else self.rows[self.header_index + 1:]


def _detect_header(rows: list[tuple[int, list[Any]]]) -> int | None:
    for i, (_, cells) in enumerate(rows[:HEADER_SCAN_ROWS]):
        if _is_header_row(cells):
            return i
    if rows:
        first = rows[0][1]
        if all(c is None or isinstance(c, str) for c in first) and not any(_looks_like_date(c) for c in first):
            return 0
    return None


class _Source:
    """The stored original file, parsed lazily per sheet."""

    def __init__(self, content: bytes, file_type: str, *, encoding: str | None = None,
                 delimiter: str | None = None):
        self.content = content
        self.file_type = file_type
        self.encoding = encoding
        self.delimiter = delimiter
        self._csv_rows: list[tuple[int, list[Any]]] | None = None
        if file_type == "csv":
            self._csv_rows, self.encoding, self.delimiter = _read_csv(content, encoding, delimiter)
            self.sheets: list[str] = []
        else:
            self.sheets = _xlsx_sheet_names(content)

    def default_sheet(self) -> str | None:
        return self.sheets[0] if self.sheets else None

    def table(self, sheet: str | None = None, header_row: int | None = None) -> _Table:
        if self.file_type == "csv":
            rows = self._csv_rows or []
            sheet = None
        else:
            sheet = sheet or self.default_sheet()
            rows = _read_xlsx(self.content, sheet)
        if header_row is None:
            header_index = _detect_header(rows)
        elif header_row == 0:
            header_index = None
        else:
            positions = [i for i, (no, _) in enumerate(rows) if no == header_row]
            if not positions:
                raise ImportValidationError(f"Zeile {header_row} ist leer oder existiert nicht.")
            header_index = positions[0]
        return _Table(sheet, rows, header_index)


# --------------------------------------------------------------------------- mapping


def guess_mapping(table: _Table) -> dict[str, int | None]:
    """Guess ``{field: column index}`` from the header names (plus the date column from the content)."""
    mapping: dict[str, int | None] = dict.fromkeys(FIELDS)
    if table.header_index is not None:
        headers = table.headers
        header_cells = table.rows[table.header_index][1]
        candidates = []
        for col, header in enumerate(headers):
            if col >= len(header_cells) or header_cells[col] is None:
                continue  # placeholder name "Spalte X"
            for f in FIELDS:
                score = _header_score(header, f)
                if score:
                    candidates.append((score, -col, f, col))
        used: set[int] = set()
        for _score, _neg, f, col in sorted(candidates, reverse=True):
            if mapping[f] is None and col not in used:
                mapping[f] = col
                used.add(col)
    if mapping["datum"] is None:
        sample = table.data[:HEADER_SCAN_ROWS]
        used = {c for c in mapping.values() if c is not None}
        for col in range(table.ncols):
            if col in used:
                continue
            values = [cells[col] for _, cells in sample if col < len(cells) and cells[col] is not None]
            if values and sum(_looks_like_date(v) for v in values) >= 0.8 * len(values):
                mapping["datum"] = col
                break
    return mapping


def _resolve_column(value: Any, table: _Table, field_name: str) -> int | None:
    if value is None or value == "":
        return None
    label = FIELD_LABELS[field_name]
    where = f" im Blatt „{table.sheet}“" if table.sheet else ""
    if isinstance(value, bool):
        raise ImportValidationError(f"Ungültige Spaltenangabe für {label}.")
    if isinstance(value, int):
        if not 0 <= value < max(table.ncols, 1):
            raise ImportValidationError(f"Spalte {value + 1} für {label} existiert nicht{where}.")
        return value
    text = str(value).strip()
    norm = _norm_header(text)
    if table.header_index is not None:
        for i, header in enumerate(table.headers):
            if _norm_header(header) == norm:
                return i
    if _COLUMN_LETTER_RE.match(text):
        index = _letter_index(text)
        if index < table.ncols:
            return index
    if text.isdigit() and int(text) < table.ncols:
        return int(text)
    raise ImportValidationError(f"Spalte „{text}“ für {label} nicht gefunden{where}.")


def _resolve_mapping(mapping: Mapping[str, Any], table: _Table, *, has_default_category: bool
                     ) -> dict[str, int | None]:
    if not isinstance(mapping, Mapping):
        raise ImportValidationError("Ungültige Spaltenzuordnung.")
    unknown = set(mapping) - set(FIELDS)
    if unknown:
        raise ImportValidationError(f"Unbekannte Felder in der Zuordnung: {', '.join(sorted(unknown))}")
    cols = {f: _resolve_column(mapping.get(f), table, f) for f in FIELDS}
    if cols["datum"] is None:
        raise ImportValidationError("Bitte die Spalte Datum zuordnen.")
    if cols["dauer"] is None and (cols["von"] is None or cols["bis"] is None):
        raise ImportValidationError("Bitte Von und Bis oder die Spalte Dauer zuordnen.")
    if cols["kategorie"] is None and not has_default_category:
        raise ImportValidationError("Bitte die Spalte Kategorie zuordnen oder eine Standard-Kategorie wählen.")
    time_cols = [cols[f] for f in TIME_FIELDS if cols[f] is not None]
    if len(time_cols) != len(set(time_cols)):
        raise ImportValidationError("Datum, Von, Bis und Dauer müssen unterschiedliche Spalten sein.")
    return cols


# --------------------------------------------------------------------------- options


@dataclass
class ImportOptions:
    """Options for validate/commit (all optional).

    ``default_category``: catalogue code for rows without category text. ``category_map``: free text →
    code for this import (remembered on commit unless ``remember_mappings`` is false).
    ``skip_errors``: commit skips error rows (False → refuse commit while errors exist).
    ``skip_duplicates``: duplicates are not imported unless listed in ``include_rows``.
    ``exclude_rows``: rows never imported. ``reason``: needed for rows in closed months.
    ``include_outside_periods``: also import rows dated outside every period (default: excluded with a
    warning; single rows can be included via ``include_rows``).
    """

    default_category: str | None = None
    category_map: dict[str, str] = field(default_factory=dict)
    remember_mappings: bool = True
    skip_errors: bool = True
    skip_duplicates: bool = True
    include_rows: frozenset[str] = frozenset()
    exclude_rows: frozenset[str] = frozenset()
    reason: str | None = None
    include_outside_periods: bool = False

    @classmethod
    def parse(cls, raw: Mapping[str, Any] | ImportOptions | None) -> ImportOptions:
        if raw is None:
            return cls()
        if isinstance(raw, ImportOptions):
            return raw
        if not isinstance(raw, Mapping):
            raise ImportValidationError("Ungültige Import-Optionen.")
        unknown = set(raw) - OPTION_KEYS
        if unknown:
            raise ImportValidationError(f"Unbekannte Optionen: {', '.join(sorted(unknown))}")
        cmap_raw = raw.get("category_map") or {}
        if not isinstance(cmap_raw, Mapping):
            raise ImportValidationError("category_map muss eine Zuordnung Text → Code sein.")
        cmap = {str(k).strip(): str(v).strip() for k, v in cmap_raw.items()
                if str(k).strip() and v is not None and str(v).strip()}

        def refs(key: str) -> frozenset[str]:
            values = raw.get(key) or []
            if isinstance(values, (str, int)):
                values = [values]
            return frozenset(str(v).strip() for v in values if str(v).strip())

        def flag(key: str, default: bool) -> bool:
            value = raw.get(key)
            return default if value is None else bool(value)

        default = raw.get("default_category")
        reason = raw.get("reason")
        return cls(
            default_category=str(default).strip() or None if default is not None else None,
            category_map=cmap,
            remember_mappings=flag("remember_mappings", True),
            skip_errors=flag("skip_errors", True),
            skip_duplicates=flag("skip_duplicates", True),
            include_rows=refs("include_rows"),
            exclude_rows=refs("exclude_rows"),
            reason=(str(reason).strip() or None) if reason is not None else None,
            include_outside_periods=flag("include_outside_periods", False),
        )

    def to_json(self) -> dict[str, Any]:
        return {
            "default_category": self.default_category,
            "category_map": dict(self.category_map),
            "remember_mappings": self.remember_mappings,
            "skip_errors": self.skip_errors,
            "skip_duplicates": self.skip_duplicates,
            "include_rows": sorted(self.include_rows),
            "exclude_rows": sorted(self.exclude_rows),
            "reason": self.reason,
            "include_outside_periods": self.include_outside_periods,
        }


def _selected(row: _Row, refs: frozenset[str]) -> bool:
    return row.ref in refs or str(row.row_no) in refs


# --------------------------------------------------------------------------- categories


class _CategoryResolver:
    def __init__(self, db: Session, options: ImportOptions):
        cats = list(db.scalars(select(Category).order_by(Category.parent_id.is_not(None), Category.sort)))
        self.by_id = {c.id: c for c in cats}
        self.by_code = {c.code.casefold(): c for c in cats}
        names: dict[str, Category | None] = {}
        for c in cats:  # main categories first (order_by), subcategory names only if unambiguous
            key = _norm_key(c.name)
            if key not in names:
                names[key] = c
            elif names[key] is not None and names[key].parent_id is not None and c.parent_id is not None:
                names[key] = None
        # Parts of main category names as aliases ("Korrektur", "Konferenzen", "Aufsichten" ...).
        aliases: dict[str, Category | None] = {}
        for c in cats:
            if c.parent_id is not None:
                continue
            for part in re.split(r"\s*(?:&|,| und )\s*", c.name):
                key = _norm_key(part)
                if len(key) > 3 and key not in names:
                    aliases[key] = None if key in aliases and aliases[key] is not c else c
        names.update({k: v for k, v in aliases.items() if v is not None})
        self.by_name = {k: v for k, v in names.items() if v is not None}
        self.remembered: dict[str, str] = {}
        for m in db.scalars(select(CategoryMapping)):
            self.remembered[_norm_key(m.source_text)] = m.category_code
        self.user: dict[str, str] = {}
        for text, code in options.category_map.items():
            if code.casefold() not in self.by_code:
                raise ImportValidationError(f"Unbekannter Kategorie-Code „{code}“ für „{text}“.")
            self.user[_norm_key(text)] = code
        self.default: Category | None = None
        if options.default_category:
            self.default = self.by_code.get(options.default_category.casefold())
            if self.default is None:
                raise ImportValidationError(f"Unbekannte Standard-Kategorie „{options.default_category}“.")

    def resolve(self, text: str | None) -> tuple[Category | None, str | None]:
        if text is None:
            return (self.default, "standard") if self.default is not None else (None, None)
        key = _norm_key(text)
        if key in self.user:
            return self.by_code[self.user[key].casefold()], "manuell"
        code = self.remembered.get(key)
        if code and code.casefold() in self.by_code:
            return self.by_code[code.casefold()], "gemerkt"
        cat = self.by_code.get(text.strip().casefold()) or self.by_name.get(key)
        if cat is not None:
            return cat, "katalog"
        return None, None

    def main_of(self, cat: Category) -> Category:
        return self.by_id.get(cat.parent_id, cat) if cat.parent_id is not None else cat


def catalogue(db: Session) -> list[dict[str, Any]]:
    """Category catalogue for the mapping UI (code, name, parent code, active)."""
    cats = list(db.scalars(select(Category).order_by(Category.parent_id.is_not(None), Category.sort)))
    by_id = {c.id: c for c in cats}
    mains = [c for c in cats if c.parent_id is None]
    result = []
    for main in sorted(mains, key=lambda c: c.sort):
        result.append({"code": main.code, "name": main.name, "parent_code": None, "active": main.active,
                       "label": f"{main.code} – {main.name}"})
        for sub in sorted((c for c in cats if c.parent_id == main.id), key=lambda c: c.sort):
            result.append({"code": sub.code, "name": sub.name, "parent_code": by_id[sub.parent_id].code,
                           "active": sub.active and main.active, "label": f"{main.code} › {sub.name}"})
    return result


# --------------------------------------------------------------------------- row analysis


@dataclass
class _Row:
    ref: str
    sheet: str | None
    row_no: int
    status: str = STATUS_OK
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    day: date | None = None
    start_at: datetime | None = None
    end_at: datetime | None = None
    duration: int | None = None
    has_bounds: bool = False
    category: Category | None = None
    main_category: Category | None = None
    category_text: str | None = None
    category_source: str | None = None
    klasse: str | None = None
    fach: str | None = None
    note: str | None = None
    location: str | None = None
    duplicate_of: str | None = None
    closed_month: bool = False
    outside_period: bool = False
    do_import: bool = False
    excluded: bool = False

    def fail(self, message: str) -> None:
        self.errors.append(message)
        self.status = STATUS_ERROR

    def to_dict(self) -> dict[str, Any]:
        cat = self.category
        return {
            "ref": self.ref,
            "sheet": self.sheet,
            "row_no": self.row_no,
            "status": self.status,
            "import": self.do_import,
            "excluded": self.excluded,
            "errors": list(self.errors),
            "warnings": list(self.warnings),
            "duplicate_of": self.duplicate_of,
            "day": self.day.isoformat() if self.day else None,
            "von": _hm_local(self.start_at) if self.has_bounds else None,
            "bis": _hm_local(self.end_at) if self.has_bounds else None,
            "duration_seconds": self.duration,
            "duration_hm": fmt_hm(self.duration) if self.duration is not None else None,
            "has_bounds": self.has_bounds,
            "category_text": self.category_text,
            "category_code": cat.code if cat else None,
            "category_name": cat.name if cat else None,
            "category_source": self.category_source,
            "klasse": self.klasse,
            "fach": self.fach,
            "note": self.note,
            "location": self.location,
            "closed_month": self.closed_month,
            "outside_period": self.outside_period,
        }


def _parse_day(value: Any) -> date:
    if isinstance(value, str):
        value = _WEEKDAY_PREFIX_RE.sub("", value.strip())
    return parse_date(value)


def _is_sum_row(cells: list[Any], cols: dict[str, int | None]) -> bool:
    """A row without a valid date whose first/date/category cell says 'Summe', 'Gesamt', ..."""
    datum = cells[cols["datum"]] if cols["datum"] is not None and cols["datum"] < len(cells) else None
    if datum is not None and _looks_like_date(datum):
        return False
    positions = {0, cols["datum"]} | ({cols["kategorie"]} if cols["kategorie"] is not None else set())
    for pos in positions:
        if pos is not None and pos < len(cells) and isinstance(cells[pos], str):
            if _SUM_WORDS_RE.search(_norm_header(cells[pos])):
                return True
    return False


@dataclass
class _Context:
    db: Session
    resolver: _CategoryResolver
    options: ImportOptions
    klassen: list[str]
    today: date
    periods: list[tuple[date, date]] = field(default_factory=list)
    _closed: dict[str, bool] = field(default_factory=dict)

    def in_period(self, d: date) -> bool:
        return any(start <= d <= end for start, end in self.periods)

    def month_closed(self, d: date) -> bool:
        ym = _ym(d)
        if ym not in self._closed:
            self._closed[ym] = self.db.get(MonthClose, ym) is not None
        return self._closed[ym]


def _parse_times(row: _Row, cells: list[Any], cols: dict[str, int | None]) -> None:
    def cell(f: str) -> Any:
        idx = cols[f]
        return cells[idx] if idx is not None and idx < len(cells) else None

    raw_day, raw_von, raw_bis, raw_dauer = cell("datum"), cell("von"), cell("bis"), cell("dauer")
    if isinstance(raw_dauer, datetime):  # Excel duration with a date-time format: offset from the epoch
        raw_dauer = raw_dauer - EXCEL_EPOCH_DT
    duration: int | None = None
    if raw_dauer is not None:
        try:
            duration = parse_duration(raw_dauer)
        except ValueError as exc:
            row.fail(f"Dauer: {exc}")
    if raw_von is None and raw_bis is None and (raw_dauer is None or duration == 0):
        row.status = STATUS_EMPTY  # e.g. weekend rows, comment lines
        row.warnings.append("Keine Zeitangabe – Zeile wird übersprungen.")
        return

    if raw_day is None:
        row.fail("Datum fehlt.")
    else:
        try:
            row.day = _parse_day(raw_day)
        except ValueError as exc:
            row.fail(str(exc))
    von = bis = None
    bis_next_day = False
    if isinstance(raw_bis, str) and raw_bis.strip()[:1].isdigit() and _NEXT_DAY_RE.search(raw_bis):
        raw_bis, bis_next_day = _NEXT_DAY_RE.sub("", raw_bis), True
    elif isinstance(raw_bis, datetime) and isinstance(raw_day, (date, datetime)):
        # XLSX date-time in "Bis": a later date than the row's date is explicit next-day notation.
        row_day = raw_day.date() if isinstance(raw_day, datetime) else raw_day
        bis_next_day = raw_bis.date() == row_day + timedelta(days=1)
    if raw_von is not None:
        try:
            von = parse_time_seconds(raw_von)
        except ValueError as exc:
            row.fail(f"Von: {exc}")
    if raw_bis is not None:
        try:
            bis = parse_time_seconds(raw_bis, end_of_day=True)
        except ValueError as exc:
            row.fail(f"Bis: {exc}")
    if row.status == STATUS_ERROR or row.day is None:
        return

    day = row.day
    if bis_next_day and von is None:
        row.fail("Bis mit „+1“ (Folgetag) braucht eine Von-Zeit.")
        return
    if von is not None and bis is not None:
        overnight = bis_next_day or bis < von
        if bis_next_day and bis >= MAX_DAY_SECONDS:
            row.fail("Bis „24:00+1“ ist keine gültige Uhrzeit.")
            return
        if von == bis and not bis_next_day:
            row.fail("Von und Bis sind gleich.")
            return
        start = _local_dt(day, von)
        end = _local_dt(day + timedelta(days=1), bis) if overnight else _local_dt(day, bis)
        wall = bis + MAX_DAY_SECONDS - von if overnight else bis - von
        if duration is not None and abs(duration - wall) > DURATION_TOLERANCE_SECONDS:
            row.fail(f"Dauer {fmt_hm(duration)} passt nicht zu Von–Bis ({fmt_hm(wall)}). "
                     "Bitte Von/Bis oder Dauer in der Zuordnung abwählen.")
            return
        if overnight and not bis_next_day and duration is None:
            # "10:00–09:00" is far more often a typo than 23 h of work: never import that silently.
            b, v = f"{bis // 3600:02d}:{bis % 3600 // 60:02d}", f"{von // 3600:02d}:{von % 3600 // 60:02d}"
            row.fail(f"Bis ({b}) liegt vor Von ({v}). Über Mitternacht bitte als Folgetag "
                     f"angeben (z. B. „{b}+1“) oder 24:00 für das Tagesende.")
            return
        if overnight:
            row.warnings.append("Über Mitternacht – wird beim Import an Mitternacht geteilt.")
        row.start_at, row.end_at, row.has_bounds = start, end, True
    elif von is not None or bis is not None:
        if duration is None:
            row.fail("Bis oder Dauer fehlt." if von is not None else "Von oder Dauer fehlt.")
            return
        if duration <= 0:
            row.fail("Die Dauer muss größer als 0 sein.")
            return
        if von is not None:
            row.start_at = _local_dt(day, von)
            row.end_at = row.start_at + timedelta(seconds=duration)
        else:
            row.end_at = _local_dt(day, bis)
            row.start_at = row.end_at - timedelta(seconds=duration)
            if local_date(row.start_at) != day:
                row.fail("Beginn (Bis minus Dauer) läge am Vortag.")
                return
        row.has_bounds = True
    else:
        row.has_bounds = False

    row.duration = int((row.end_at - row.start_at).total_seconds()) if row.has_bounds else duration
    if row.duration is None or row.duration <= 0:
        row.fail("Die Dauer muss größer als 0 sein.")
    elif row.duration > MAX_DAY_SECONDS:
        row.fail("Ein Eintrag darf höchstens 24 Stunden dauern.")


def _parse_attrs(row: _Row, cells: list[Any], cols: dict[str, int | None], ctx: _Context) -> None:
    def text(f: str) -> str | None:
        idx = cols[f]
        return _cell_text(cells[idx]) if idx is not None and idx < len(cells) else None

    extras: list[str] = []
    row.category_text = text("kategorie")
    cat, source = ctx.resolver.resolve(row.category_text)
    if cat is None:
        row.fail(f"Kategorie „{row.category_text}“ ist keinem Code zugeordnet." if row.category_text
                 else "Kategorie fehlt.")
    else:
        main = ctx.resolver.main_of(cat)
        row.category, row.main_category, row.category_source = cat, main, source
        if not cat.active or not main.active:
            row.fail(f"Kategorie „{cat.name}“ ist deaktiviert.")

    klasse = text("klasse")
    if klasse:
        if ctx.klassen:
            match = next((k for k in ctx.klassen if k.casefold() == klasse.casefold()), None)
            if match is None:
                row.warnings.append(f"Klasse „{klasse}“ ist nicht in den Einstellungen – wird in die Notiz übernommen.")
                extras.append(f"Klasse {klasse}")
            else:
                row.klasse = match
        else:
            row.klasse = klasse
    row.fach = text("fach")

    ort = text("ort")
    if ort:
        location = LOCATION_SYNONYMS.get(_norm_header(ort))
        if location is None:
            row.warnings.append(f"Ort „{ort}“ unbekannt – wird in die Notiz übernommen.")
            extras.append(f"Ort: {ort}")
        row.location = location

    note = text("notiz")
    parts = [p for p in (note, *extras) if p]
    row.note = " · ".join(parts) if parts else None
    if row.main_category is not None and row.main_category.requires_note and not row.note:
        row.fail(f"Für die Kategorie „{row.main_category.name}“ ist eine Notiz Pflicht.")

    if row.day is not None:
        if row.day > ctx.today:
            row.warnings.append("Datum liegt in der Zukunft.")
        if not ctx.in_period(row.day):
            row.outside_period = True
            row.warnings.append(
                f"{OUTSIDE_PERIOD_WARNING} ({row.day:%d.%m.%Y}) – "
                + ("wird trotzdem importiert." if ctx.options.include_outside_periods
                   else "wird nicht importiert (Option „auch außerhalb importieren“ oder Zeile einschließen).")
            )
        if ctx.month_closed(row.day):
            row.closed_month = True
            label = row.day.strftime("%m/%Y")
            if ctx.options.reason:
                row.warnings.append(f"Monat {label} ist abgeschlossen – Import als Korrektur mit Begründung.")
            else:
                row.fail(f"Monat {label} ist abgeschlossen – Import nur mit Begründung.")


def _parse_row(row_no: int, cells: list[Any], table: _Table, cols: dict[str, int | None], ctx: _Context) -> _Row:
    ref = f"{table.sheet}!{row_no}" if table.sheet is not None else str(row_no)
    row = _Row(ref=ref, sheet=table.sheet, row_no=row_no)
    if _is_sum_row(cells, cols):
        row.status = STATUS_IGNORED
        row.warnings.append("Summenzeile – wird nicht importiert.")
        return row
    _parse_times(row, cells, cols)
    if row.status == STATUS_EMPTY:
        return row
    _parse_attrs(row, cells, cols, ctx)
    return row


def _entry_label(e: TimeEntry) -> str:
    code = e.category.code if e.category else "?"
    when = e.day.strftime("%d.%m.%Y")
    if e.has_bounds:
        end = _hm_local(e.end_at) if e.end_at else "läuft"
        when += f" {_hm_local(e.start_at)}–{end}"
    else:
        when += f" ({fmt_hm(e.duration_seconds)})"
    return f"{when}, {code}, Quelle {e.source}"


def _load_existing(db: Session, rows: list[_Row]) -> list[TimeEntry]:
    days = [r.day for r in rows if r.day is not None]
    if not days:
        return []
    lo, hi = min(days) - timedelta(days=1), max(days) + timedelta(days=1)
    return list(db.scalars(
        select(TimeEntry).where(
            TimeEntry.deleted_at.is_(None),
            TimeEntry.status == EntryStatus.ACTIVE.value,
            or_(TimeEntry.day.between(lo, hi), TimeEntry.end_at.is_(None)),
        ).order_by(TimeEntry.start_at)
    ))


def _same_entry(a_bounded: bool, a_start: datetime | None, b_bounded: bool, b_start: datetime | None) -> bool:
    """Duplicate rule: same day/duration/category; two bounded items must also start together."""
    return not (a_bounded and b_bounded) or a_start == b_start


def _mark_duplicates(rows: list[_Row], existing: list[TimeEntry]) -> None:
    by_key: dict[tuple, list[TimeEntry]] = defaultdict(list)
    for e in existing:
        if e.end_at is not None:
            by_key[(e.day, e.duration_seconds, e.category_id)].append(e)
    seen: dict[tuple, list[_Row]] = defaultdict(list)
    for row in rows:
        if row.status != STATUS_OK:
            continue
        key = (row.day, row.duration, row.main_category.id)
        hit = next((e for e in by_key.get(key, []) if _same_entry(row.has_bounds, row.start_at, e.has_bounds,
                                                                     e.start_at)), None)
        if hit is not None:
            row.status = STATUS_DUPLICATE
            row.duplicate_of = f"eintrag:{hit.id}"
            row.warnings.append(f"Doppelt: gleicher Tag, gleiche Dauer und Kategorie wie bestehender Eintrag "
                                f"({_entry_label(hit)}).")
            continue
        earlier = next((r for r in seen.get(key, []) if _same_entry(row.has_bounds, row.start_at, r.has_bounds,
                                                                      r.start_at)), None)
        if earlier is not None:
            row.status = STATUS_DUPLICATE
            row.duplicate_of = earlier.ref
            row.warnings.append(f"Doppelt: gleicher Tag, gleiche Dauer und Kategorie wie Zeile {earlier.ref}.")
            continue
        seen[key].append(row)


def _decide_import(rows: list[_Row], options: ImportOptions) -> None:
    for row in rows:
        row.excluded = _selected(row, options.exclude_rows)
        if row.excluded:
            row.do_import = False
        elif row.status == STATUS_OK:
            row.do_import = True
        elif row.status == STATUS_DUPLICATE:
            row.do_import = not options.skip_duplicates or _selected(row, options.include_rows)
        else:
            row.do_import = False
        if row.do_import and row.outside_period and not options.include_outside_periods:
            row.do_import = _selected(row, options.include_rows)


def _touched_days(start: datetime, end: datetime) -> Iterable[date]:
    d, last = local_date(start), local_date(end - timedelta(seconds=1)) if end > start else local_date(start)
    while d <= last:
        yield d
        d += timedelta(days=1)


def _mark_overlaps(rows: list[_Row], existing: list[TimeEntry]) -> None:
    """Bounded rows to be imported must not overlap bounded entries or each other (half-open)."""
    buckets: dict[date, list[tuple[datetime, datetime, str]]] = defaultdict(list)
    running: list[tuple[datetime, str]] = []
    for e in existing:
        if not e.has_bounds:
            continue
        label = f"bestehendem Eintrag ({_entry_label(e)})"
        if e.end_at is None:
            running.append((e.start_at, "dem laufenden Timer"))
            continue
        for d in _touched_days(e.start_at, e.end_at):
            buckets[d].append((e.start_at, e.end_at, label))
    for row in rows:
        if not (row.do_import and row.has_bounds):
            continue
        start, end = row.start_at, row.end_at
        hit = next((label for s, label in running if end > s), None)
        if hit is None:
            for d in _touched_days(start, end):
                hit = next((label for s, e, label in buckets.get(d, []) if s < end and e > start), None)
                if hit:
                    break
        if hit is not None:
            row.fail(f"Überschneidung mit {hit}.")
            row.do_import = False
            continue
        for d in _touched_days(start, end):
            buckets[d].append((start, end, f"Zeile {row.ref}"))


@dataclass
class _Analysis:
    batch: ImportBatch
    tables: list[_Table]
    cols: dict[str | None, dict[str, int | None]]
    options: ImportOptions
    rows: list[_Row]
    existing_seconds: dict[str, int]

    def counts(self) -> dict[str, int]:
        c = Counter(r.status for r in self.rows)
        return {
            "rows_total": len(self.rows),
            STATUS_OK: c[STATUS_OK],
            STATUS_ERROR: c[STATUS_ERROR],
            STATUS_DUPLICATE: c[STATUS_DUPLICATE],
            STATUS_EMPTY: c[STATUS_EMPTY],
            STATUS_IGNORED: c[STATUS_IGNORED],
            "ausgeschlossen": sum(1 for r in self.rows if r.excluded),
            "ausserhalb": sum(1 for r in self.rows
                              if r.outside_period and r.status in (STATUS_OK, STATUS_DUPLICATE)),
            "warnungen": sum(1 for r in self.rows if r.warnings),
            "importieren": sum(1 for r in self.rows if r.do_import),
        }


def _analyse(
    db: Session,
    batch: ImportBatch,
    mapping: Mapping[str, Any],
    *,
    sheets: str | list[str] | None,
    header_row: int | None,
    options: ImportOptions,
    today: date | None,
) -> _Analysis:
    source = _load_source(batch)
    if isinstance(sheets, str):
        sheets = [sheets]
    if source.file_type == "csv":
        tables = [source.table(None, header_row)]
    else:
        names = list(dict.fromkeys(sheets or [source.default_sheet()]))
        tables = [source.table(name, header_row) for name in names]

    resolver = _CategoryResolver(db, options)
    settings = db.get(Settings, 1)
    ctx = _Context(db=db, resolver=resolver, options=options,
                   klassen=list(settings.klassen or []) if settings else [], today=today or today_local(),
                   periods=[(p.start, p.end) for p in db.scalars(select(Period))])
    rows: list[_Row] = []
    cols_by_sheet: dict[str | None, dict[str, int | None]] = {}
    for table in tables:
        cols = _resolve_mapping(mapping, table, has_default_category=resolver.default is not None)
        cols_by_sheet[table.sheet] = cols
        for row_no, cells in table.data:
            rows.append(_parse_row(row_no, cells, table, cols, ctx))

    existing = _load_existing(db, rows)
    _mark_duplicates(rows, existing)
    _decide_import(rows, options)
    _mark_overlaps(rows, existing)

    months = {_ym(r.day) for r in rows if r.day is not None}
    existing_seconds: dict[str, int] = {}
    if months:
        lo = min(date.fromisoformat(m + "-01") for m in months)
        hi_month = max(months)
        y, m = (int(x) for x in hi_month.split("-"))
        hi = date(y + (m == 12), m % 12 + 1, 1)
        q = (select(TimeEntry.day, TimeEntry.duration_seconds)
             .where(TimeEntry.deleted_at.is_(None), TimeEntry.status == EntryStatus.ACTIVE.value,
                    TimeEntry.day >= lo, TimeEntry.day < hi))
        for d, secs in db.execute(q):
            ym = _ym(d)
            if ym in months:
                existing_seconds[ym] = existing_seconds.get(ym, 0) + (secs or 0)
    return _Analysis(batch, tables, cols_by_sheet, options, rows, existing_seconds)


def _report(analysis: _Analysis) -> dict[str, Any]:
    rows = analysis.rows
    counts = analysis.counts()

    months: dict[str, dict[str, Any]] = {}
    for r in rows:
        if r.day is None or r.duration is None or r.status in (STATUS_EMPTY, STATUS_IGNORED):
            continue
        ym = _ym(r.day)
        m = months.setdefault(ym, {"month": ym, "label": _month_label(ym), "rows": 0, "rows_import": 0,
                                   "datei_seconds": 0, "import_seconds": 0, "fehler": 0, "duplikate": 0})
        m["rows"] += 1
        m["datei_seconds"] += r.duration
        if r.do_import:
            m["rows_import"] += 1
            m["import_seconds"] += r.duration
        if r.status == STATUS_ERROR:
            m["fehler"] += 1
        if r.status == STATUS_DUPLICATE:
            m["duplikate"] += 1
    month_list = []
    for ym in sorted(months):
        m = months[ym]
        m["differenz_seconds"] = m["import_seconds"] - m["datei_seconds"]
        m["bestehend_seconds"] = analysis.existing_seconds.get(ym, 0)
        for key in ("datei", "import", "differenz", "bestehend"):
            m[f"{key}_hm"] = fmt_hm(m[f"{key}_seconds"])
        month_list.append(m)

    unmapped: dict[str, dict[str, Any]] = {}
    resolution: dict[str, dict[str, Any]] = {}
    for r in rows:
        if r.category_text is None:
            continue
        key = _norm_key(r.category_text)
        if r.category is None and r.status == STATUS_ERROR:
            u = unmapped.setdefault(key, {"text": r.category_text, "count": 0, "rows": []})
            u["count"] += 1
            if len(u["rows"]) < 10:
                u["rows"].append(r.ref)
        elif r.category is not None:
            res = resolution.setdefault(key, {"text": r.category_text, "code": r.category.code,
                                              "name": r.category.name, "source": r.category_source, "count": 0})
            res["count"] += 1

    seconds_import = sum(r.duration or 0 for r in rows if r.do_import)
    seconds_file = sum(m["datei_seconds"] for m in month_list)
    summary = {**counts, "seconds_import": seconds_import, "import_hm": fmt_hm(seconds_import),
               "seconds_file": seconds_file, "file_hm": fmt_hm(seconds_file),
               "entries_bounded": sum(1 for r in rows if r.do_import and r.has_bounds),
               "entries_unbounded": sum(1 for r in rows if r.do_import and not r.has_bounds)}
    closed = sorted({_ym(r.day) for r in rows if r.closed_month and r.day is not None})

    hinweise = []
    if counts[STATUS_ERROR]:
        hinweise.append(f"{counts[STATUS_ERROR]} Zeile(n) mit Fehlern werden nicht importiert.")
    if unmapped:
        hinweise.append(f"{len(unmapped)} Kategorie-Text(e) ohne Zuordnung – bitte einem Code zuordnen.")
    if counts[STATUS_DUPLICATE]:
        verb = "werden übersprungen" if analysis.options.skip_duplicates else "werden trotzdem importiert"
        hinweise.append(f"{counts[STATUS_DUPLICATE]} mögliche Duplikate {verb}.")
    if closed:
        hinweise.append("Zeilen in abgeschlossenen Monaten: " + ", ".join(_month_label(m) for m in closed)
                        + (" – Import als Korrektur mit Begründung." if analysis.options.reason
                           else " – Import nur mit Begründung."))
    if counts["ausserhalb"]:
        hinweise.append(
            f"{counts['ausserhalb']} Zeile(n) liegen außerhalb jedes Schuljahres"
            + (" und werden trotzdem importiert." if analysis.options.include_outside_periods
               else " und werden nicht importiert. Schuljahr anlegen oder Option „auch außerhalb importieren“ "
                    "wählen.")
        )
    if summary["entries_unbounded"]:
        hinweise.append("Zeilen nur mit Dauer werden ohne Uhrzeiten (Kennzeichen „ohne Zeitgrenzen“) übernommen.")

    can_commit = counts["importieren"] > 0 and (analysis.options.skip_errors or counts[STATUS_ERROR] == 0)
    return {
        "batch_id": analysis.batch.id,
        "stage": _state(analysis.batch).get("stage"),
        "sheets": [t.sheet for t in analysis.tables if t.sheet is not None],
        "header_rows": {(t.sheet or ""): t.header_row for t in analysis.tables},
        "columns": {(sheet or ""): cols for sheet, cols in analysis.cols.items()},
        "options": analysis.options.to_json(),
        "summary": summary,
        "rows": [r.to_dict() for r in rows],
        "unmapped_categories": sorted(unmapped.values(), key=lambda u: (-u["count"], u["text"].casefold())),
        "category_resolution": sorted(resolution.values(), key=lambda u: (-u["count"], u["text"].casefold())),
        "months": month_list,
        "closed_months": closed,
        "hinweise": hinweise,
        "can_commit": can_commit,
    }


# --------------------------------------------------------------------------- batches


def _state(batch: ImportBatch) -> dict[str, Any]:
    return dict(batch.mapping_json or {})


def _set_state(batch: ImportBatch, **changes: Any) -> None:
    state = _state(batch)
    state.update(changes)
    batch.mapping_json = state  # new object → SQLAlchemy detects the JSON change


_BATCH_ID_RE = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}")
_FILE_TYPE_RE = re.compile(r"[a-z0-9]{1,8}")


def _batch_dir(batch_id: str) -> Path:
    """``DATA_DIR/imports/<id>``; the id must be a UUID (file handling never leaves the imports folder)."""
    if not _BATCH_ID_RE.fullmatch(str(batch_id or "")):
        raise ImportNotFound("Import nicht gefunden.")
    return get_config().imports_dir / batch_id


def _stored_path(batch: ImportBatch) -> Path:
    """The stored original file – always ``imports/<id>/original.<type>`` (``stored_file`` is informational)."""
    file_type = str(_state(batch).get("file_type") or "csv")
    if not _FILE_TYPE_RE.fullmatch(file_type):
        raise ImportNotFound("Die Originaldatei dieses Imports ist nicht mehr vorhanden.")
    return _batch_dir(batch.id) / f"original.{file_type}"


def _load_source(batch: ImportBatch) -> _Source:
    state = _state(batch)
    path = _stored_path(batch)
    try:
        content = path.read_bytes()
    except OSError as exc:
        raise ImportNotFound("Die Originaldatei dieses Imports ist nicht mehr vorhanden.") from exc
    return _Source(content, state.get("file_type", "csv"), encoding=state.get("encoding"),
                   delimiter=state.get("delimiter"))


def get_batch(db: Session, batch_id: str) -> ImportBatch:
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise ImportNotFound("Import nicht gefunden.")
    return batch


def batch_label(batch: ImportBatch) -> str:
    """Report marker (Konzept §12): 'Import vom DD.MM.YYYY aus Datei X'."""
    return f"Import vom {to_local(batch.created_at):%d.%m.%Y} aus Datei {batch.filename}"


def _entry_counts(db: Session, batch_ids: list[str]) -> dict[str, tuple[int, int]]:
    """{batch_id: (active entries, entries in closed months / locked)}."""
    if not batch_ids:
        return {}
    active: dict[str, int] = defaultdict(int)
    locked: dict[str, int] = defaultdict(int)
    q = (select(TimeEntry.import_batch_id, TimeEntry.day, TimeEntry.locked)
         .where(TimeEntry.import_batch_id.in_(batch_ids), TimeEntry.deleted_at.is_(None)))
    closed_cache: dict[str, bool] = {}
    for bid, d, is_locked in db.execute(q):
        active[bid] += 1
        ym = _ym(d)
        if ym not in closed_cache:
            closed_cache[ym] = db.get(MonthClose, ym) is not None
        if is_locked or closed_cache[ym]:
            locked[bid] += 1
    return {bid: (active[bid], locked[bid]) for bid in batch_ids}


def batch_to_dict(db: Session, batch: ImportBatch, *, counts: tuple[int, int] | None = None) -> dict[str, Any]:
    state = _state(batch)
    if counts is None:
        counts = _entry_counts(db, [batch.id]).get(batch.id, (0, 0))
    return {
        "id": batch.id,
        "filename": batch.filename,
        "label": batch_label(batch),
        "created_at": _iso_utc(batch.created_at),
        "created_local": to_local(batch.created_at).strftime("%d.%m.%Y %H:%M"),
        "stage": state.get("stage"),
        "file_type": state.get("file_type"),
        "encoding": state.get("encoding"),
        "delimiter": state.get("delimiter"),
        "sheets": state.get("sheets") or [],
        "selected_sheets": state.get("selected_sheets") or [],
        "mapping": state.get("mapping"),
        "rows_total": batch.rows_total,
        "rows_imported": batch.rows_imported,
        "rows_skipped": batch.rows_skipped,
        "entries_created": state.get("entries_created", 0),
        "entries_active": counts[0],
        "entries_locked": counts[1],
        "committed_at": state.get("committed_at"),
        "undone_at": _iso_utc(batch.undone_at),
        "notes": batch.notes,
    }


def list_batches(db: Session) -> list[dict[str, Any]]:
    """All batches, newest first."""
    batches = list(db.scalars(select(ImportBatch).order_by(ImportBatch.created_at.desc(), ImportBatch.id.desc())))
    counts = _entry_counts(db, [b.id for b in batches])
    return [batch_to_dict(db, b, counts=counts.get(b.id, (0, 0))) for b in batches]


# --------------------------------------------------------------------------- public API: preview


def _preview_dict(db: Session, batch: ImportBatch, source: _Source, table: _Table) -> dict[str, Any]:
    state = _state(batch)
    headers = table.headers
    data = table.data
    columns = []
    for i, header in enumerate(headers):
        samples = [display_cell(cells[i]) for _, cells in data[:PREVIEW_ROWS]
                   if i < len(cells) and cells[i] is not None][:3]
        columns.append({"index": i, "letter": column_letter(i), "header": header, "samples": samples})
    return {
        "batch_id": batch.id,
        "filename": batch.filename,
        "stage": state.get("stage"),
        "file_type": source.file_type,
        "encoding": source.encoding,
        "delimiter": source.delimiter,
        "sheets": source.sheets,
        "sheet": table.sheet,
        "header_row": table.header_row,
        "headers": headers,
        "columns": columns,
        "rows": [{"row_no": no, "cells": [display_cell(cells[i]) if i < len(cells) else "" for i in range(len(headers))]}
                 for no, cells in data[:PREVIEW_ROWS]],
        "rows_total": len(data),
        "mapping": guess_mapping(table),
        "fields": [{"key": f, "label": FIELD_LABELS[f], "required": f == "datum"} for f in FIELDS],
        "categories": catalogue(db),
        "remembered_mappings": len(list_mappings(db)),
    }


def upload(db: Session, content: bytes, filename: str | None, *, sheet: str | None = None,
           now: datetime | None = None) -> dict[str, Any]:
    """Store the original file, create the batch and return the preview (see module docstring)."""
    if not content:
        raise ImportFileError("Die Datei ist leer.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ImportFileTooLarge(f"Die Datei ist zu groß (höchstens {MAX_UPLOAD_BYTES // (1024 * 1024)} MB).")
    file_type = detect_file_type(content, filename)
    source = _Source(content, file_type)
    if sheet is not None and file_type == "xlsx" and sheet not in source.sheets:
        raise ImportValidationError(f"Tabellenblatt „{sheet}“ nicht gefunden.")
    table = source.table(sheet)
    if not table.rows:
        raise ImportFileError("Die Datei enthält keine Daten.")

    safe_name = Path(str(filename or "").replace("\\", "/")).name.strip()[:255] or f"import.{file_type}"
    batch_id = uuid7()
    directory = _batch_dir(batch_id)
    path = directory / f"original.{file_type}"
    directory.mkdir(parents=True, exist_ok=True)
    try:
        path.write_bytes(content)
        batch = ImportBatch(
            id=batch_id,
            created_at=to_naive(now) if now else now_utc(),
            filename=safe_name,
            stored_file=str(path),
            mapping_json={
                "stage": STAGE_UPLOADED,
                "file_type": file_type,
                "encoding": source.encoding,
                "delimiter": source.delimiter,
                "sheets": source.sheets,
                "size_bytes": len(content),
            },
            rows_total=len(table.data),
        )
        db.add(batch)
        db.commit()
    except Exception:
        db.rollback()
        shutil.rmtree(directory, ignore_errors=True)
        raise
    return _preview_dict(db, batch, source, table)


def preview(db: Session, batch_id: str, *, sheet: str | None = None, header_row: int | None = None
            ) -> dict[str, Any]:
    """Preview of a stored batch (e.g. another sheet or another header row)."""
    batch = get_batch(db, batch_id)
    source = _load_source(batch)
    if sheet is not None and source.file_type == "xlsx" and sheet not in source.sheets:
        raise ImportValidationError(f"Tabellenblatt „{sheet}“ nicht gefunden.")
    return _preview_dict(db, batch, source, source.table(sheet, header_row))


# --------------------------------------------------------------------------- public API: validate / commit


def validate(
    db: Session,
    batch_id: str,
    mapping: Mapping[str, Any],
    *,
    sheets: str | list[str] | None = None,
    header_row: int | None = None,
    options: Mapping[str, Any] | ImportOptions | None = None,
    today: date | None = None,
) -> dict[str, Any]:
    """Check every row with ``mapping`` and return the report (no entries are written).

    The mapping/options are stored on the batch so :func:`commit` can be called with the batch id alone.
    """
    batch = get_batch(db, batch_id)
    opts = ImportOptions.parse(options)
    analysis = _analyse(db, batch, mapping, sheets=sheets, header_row=header_row, options=opts, today=today)
    state = _state(batch)
    if state.get("stage") in (STAGE_UPLOADED, STAGE_VALIDATED, None):
        _set_state(batch, stage=STAGE_VALIDATED, mapping=dict(mapping),
                   selected_sheets=[t.sheet for t in analysis.tables if t.sheet is not None],
                   header_row=header_row, options=opts.to_json())
        batch.rows_total = len(analysis.rows)
        db.commit()
    return _report(analysis)


def _remember(db: Session, category_map: Mapping[str, str]) -> int:
    if not category_map:
        return 0
    existing = {_norm_key(m.source_text): m for m in db.scalars(select(CategoryMapping))}
    count = 0
    for text, code in category_map.items():
        text = text.strip()[:200]
        if not text or not code:
            continue
        row = existing.get(_norm_key(text))
        if row is None:
            row = CategoryMapping(source_text=text, category_code=code)
            db.add(row)
            existing[_norm_key(text)] = row
        else:
            row.category_code = code
        count += 1
    return count


def commit(
    db: Session,
    batch_id: str,
    mapping: Mapping[str, Any] | None = None,
    *,
    sheets: str | list[str] | None = None,
    header_row: int | None = None,
    options: Mapping[str, Any] | ImportOptions | None = None,
    now: datetime | None = None,
    today: date | None = None,
) -> dict[str, Any]:
    """Import the rows of a batch (all or nothing). Without ``mapping`` the last validated one is used."""
    batch = get_batch(db, batch_id)
    state = _state(batch)
    stage = state.get("stage")
    if stage == STAGE_COMMITTED:
        raise ImportConflict("Dieser Import wurde bereits übernommen.")
    if stage == STAGE_UNDONE or batch.undone_at is not None:
        raise ImportConflict("Dieser Import wurde rückgängig gemacht. Bitte die Datei erneut hochladen.")
    if mapping is None:
        mapping = state.get("mapping")
        if not mapping:
            raise ImportValidationError("Bitte zuerst die Spalten zuordnen.")
        if sheets is None:
            sheets = state.get("selected_sheets") or None
        if header_row is None:
            header_row = state.get("header_row")
        if options is None:
            options = state.get("options")
    opts = ImportOptions.parse(options)
    now = to_naive(now) if now is not None else now_utc().replace(microsecond=0)

    analysis = _analyse(db, batch, mapping, sheets=sheets, header_row=header_row, options=opts, today=today)
    counts = analysis.counts()
    if counts[STATUS_ERROR] and not opts.skip_errors:
        raise ImportValidationError(
            f"{counts[STATUS_ERROR]} Zeile(n) enthalten Fehler. Bitte korrigieren oder Fehlerzeilen überspringen."
        )
    to_import = [r for r in analysis.rows if r.do_import]
    if not to_import:
        raise ImportValidationError("Keine Zeilen zum Importieren.")

    created = 0
    current: _Row | None = None
    try:
        for row in to_import:
            current = row
            data: dict[str, Any] = {
                "category_id": row.category.id,
                "klasse": row.klasse,
                "fach": row.fach,
                "note": row.note,
                "location": row.location,
                "import_batch_id": batch.id,
            }
            if row.has_bounds:
                data.update(has_bounds=True, start_at=row.start_at, end_at=row.end_at)
            else:
                data.update(has_bounds=False, day=row.day, duration_seconds=row.duration)
            reason = f"{COMMIT_REASON_PREFIX}: {opts.reason}" if row.closed_month and opts.reason else None
            parts = entries_svc.create_entry(db, data, source=EntrySource.IMPORT.value, reason=reason, now=now,
                                             commit=False)
            created += len(parts)
    except entries_svc.DomainError as exc:
        db.rollback()
        where = f"Zeile {current.ref}: " if current is not None else ""
        raise ImportValidationError(f"{where}{exc.message} – es wurde nichts importiert.") from exc

    remembered = _remember(db, opts.category_map) if opts.remember_mappings else 0
    skipped = {
        "fehler": counts[STATUS_ERROR],
        "duplikat": sum(1 for r in analysis.rows if r.status == STATUS_DUPLICATE and not r.do_import),
        "leer": counts[STATUS_EMPTY],
        "ignoriert": counts[STATUS_IGNORED],
        "ausgeschlossen": counts["ausgeschlossen"],
    }
    batch.rows_total = len(analysis.rows)
    batch.rows_imported = len(to_import)
    batch.rows_skipped = batch.rows_total - batch.rows_imported
    parts = [f"{len(to_import)} Zeile(n) importiert ({created} Einträge)"]
    labels = {"duplikat": "Duplikat(e)", "fehler": "Fehlerzeile(n)", "ausgeschlossen": "ausgeschlossen",
              "leer": "ohne Zeitangabe", "ignoriert": "Summenzeile(n)"}
    parts += [f"{n} {labels[k]} übersprungen" for k, n in skipped.items() if n]
    batch.notes = ", ".join(parts) + "."
    _set_state(batch, stage=STAGE_COMMITTED, mapping=dict(mapping),
               selected_sheets=[t.sheet for t in analysis.tables if t.sheet is not None],
               header_row=header_row, options=opts.to_json(), committed_at=_iso_utc(now),
               entries_created=created, counts={**counts, "skipped": skipped})
    db.commit()
    report = _report(analysis)
    return {
        "batch": batch_to_dict(db, batch),
        "imported_rows": len(to_import),
        "entries_created": created,
        "skipped": skipped,
        "mappings_remembered": remembered,
        "summary": report["summary"],
        "months": report["months"],
    }


# --------------------------------------------------------------------------- public API: undo


def undo(db: Session, batch_id: str, *, reason: str | None = None, now: datetime | None = None) -> dict[str, Any]:
    """Soft-delete all entries of a batch (reason 'Import rückgängig'); discard a never-committed batch.

    Entries in closed months need ``reason`` – without it nothing is changed and :class:`ImportLocked`
    reports the number of affected entries.
    """
    batch = get_batch(db, batch_id)
    state = _state(batch)
    if batch.undone_at is not None or state.get("stage") == STAGE_UNDONE:
        raise ImportConflict("Dieser Import wurde bereits rückgängig gemacht.")
    reason = (reason or "").strip() or None
    now = to_naive(now) if now is not None else now_utc().replace(microsecond=0)
    entries = list(db.scalars(
        select(TimeEntry).where(TimeEntry.import_batch_id == batch.id, TimeEntry.deleted_at.is_(None))
        .order_by(TimeEntry.start_at)
    ))

    if state.get("stage") != STAGE_COMMITTED and not entries:
        has_entries = db.scalar(select(func.count(TimeEntry.id)).where(TimeEntry.import_batch_id == batch.id))
        if not has_entries:
            db.delete(batch)
            db.commit()
            shutil.rmtree(_batch_dir(batch_id), ignore_errors=True)
            return {"discarded": True, "batch_id": batch_id, "entries_deleted": 0, "locked_entries": 0}

    closed_cache: dict[str, bool] = {}

    def is_locked(e: TimeEntry) -> bool:
        ym = _ym(e.day)
        if ym not in closed_cache:
            closed_cache[ym] = db.get(MonthClose, ym) is not None
        return bool(e.locked) or closed_cache[ym]

    locked = [e for e in entries if is_locked(e)]
    if locked and not reason:
        months = ", ".join(_month_label(m) for m in sorted({_ym(e.day) for e in locked}))
        raise ImportLocked(
            f"{len(locked)} Eintrag/Einträge dieses Imports liegen in abgeschlossenen Monaten ({months}). "
            "Rückgängig nur mit Begründung.",
            locked_entries=len(locked),
        )
    delete_reason = f"{UNDO_REASON}: {reason}" if reason else UNDO_REASON
    try:
        for e in entries:
            entries_svc.delete_entry(db, e.id, reason=delete_reason, now=now, commit=False)
    except entries_svc.DomainError as exc:
        db.rollback()
        raise ImportConflict(f"{exc.message} – Import wurde nicht rückgängig gemacht.") from exc
    batch.undone_at = now
    note = f"Rückgängig gemacht am {to_local(now):%d.%m.%Y %H:%M} ({len(entries)} Einträge gelöscht)."
    batch.notes = f"{batch.notes} {note}".strip() if batch.notes else note
    _set_state(batch, stage=STAGE_UNDONE, undo_reason=reason)
    db.commit()
    return {"discarded": False, "batch": batch_to_dict(db, batch), "batch_id": batch.id,
            "entries_deleted": len(entries), "locked_entries": len(locked)}


# --------------------------------------------------------------------------- category mapping memory


def list_mappings(db: Session) -> list[dict[str, Any]]:
    cats = {c.code: c for c in db.scalars(select(Category))}
    result = []
    for m in db.scalars(select(CategoryMapping).order_by(CategoryMapping.source_text)):
        cat = cats.get(m.category_code)
        result.append({
            "source_text": m.source_text,
            "category_code": m.category_code,
            "category_name": cat.name if cat else None,
            "valid": cat is not None,
            "updated_at": _iso_utc(m.updated_at),
        })
    return result


def save_mapping(db: Session, source_text: str, category_code: str) -> dict[str, Any]:
    text = (source_text or "").strip()
    code = (category_code or "").strip()
    if not text:
        raise ImportValidationError("Text fehlt.")
    cat = db.scalar(select(Category).where(func.lower(Category.code) == code.casefold())) if code else None
    if cat is None:
        raise ImportValidationError(f"Unbekannter Kategorie-Code „{code}“.")
    _remember(db, {text: cat.code})
    db.commit()
    return next(m for m in list_mappings(db) if _norm_key(m["source_text"]) == _norm_key(text))


def delete_mapping(db: Session, source_text: str) -> None:
    key = _norm_key(source_text or "")
    rows = [m for m in db.scalars(select(CategoryMapping)) if _norm_key(m.source_text) == key]
    if not rows:
        raise ImportNotFound("Zuordnung nicht gefunden.")
    for m in rows:
        db.delete(m)
    db.commit()
